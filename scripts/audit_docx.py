#!/usr/bin/env python3
"""Audit study-fast DOCX files for fixed-template and privacy requirements."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

from common import get_template, privacy_errors_and_warnings, resolve_layout

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError as exc:
    raise SystemExit("python-docx is required: python -m pip install -r requirements.txt") from exc


ALLOWED_COLORS = {None, "000000", "auto"}


def iter_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def font_value(run, attribute):
    r_pr = run._element.rPr
    if r_pr is None or r_pr.rFonts is None:
        return None
    return r_pr.rFonts.get(qn("w:" + attribute))


def audit(path, prohibited_terms, requested_template=None, privacy_mode="balanced", document_config=None):
    errors, warnings = [], []
    try:
        document = Document(path)
    except Exception as exc:
        return ["DOCX cannot be opened: %s" % exc], warnings
    try:
        layout = resolve_layout(document_config or {"document_type": document.core_properties.subject})
    except ValueError as exc:
        return [str(exc)], warnings
    allowed_colors = ALLOWED_COLORS | {layout["question_color"]}
    embedded_template = document.core_properties.category or None
    template_name = requested_template or embedded_template
    try:
        template_name, template = get_template(template_name)
    except ValueError as exc:
        return [str(exc)], warnings
    latin_font = template["fonts"]["latin"]
    cjk_font = template["fonts"]["cjk"]
    margins = template["page"]["margins_cm"]
    for section_index, section in enumerate(document.sections):
        tolerances = {
            "page_width": (section.page_width.cm, 21.0),
            "page_height": (section.page_height.cm, 29.7),
            "top_margin": (section.top_margin.cm, margins["top"]),
            "bottom_margin": (section.bottom_margin.cm, margins["bottom"]),
            "left_margin": (section.left_margin.cm, margins["left"]),
            "right_margin": (section.right_margin.cm, margins["right"]),
        }
        for label, (actual, expected) in tolerances.items():
            if abs(actual - expected) > 0.03:
                errors.append("section %d %s is %.2f cm; expected %.2f cm." % (
                    section_index + 1, label, actual, expected
                ))
    if not layout["header"] and any(p.text for sec in document.sections for h in (sec.header, sec.first_page_header, sec.even_page_header) for p in h.paragraphs):
        errors.append("Header text is forbidden by the selected layout.")
    for table in document.tables:
        if layout["table_style"] == "three_line":
            borders = table._tbl.tblPr.find(qn("w:tblBorders"))
            if borders is None:
                errors.append("Missing three-line table borders.")
            else:
                for edge in ("left", "right", "insideV", "insideH"):
                    item = borders.find(qn("w:" + edge))
                    if item is None or item.get(qn("w:val")) not in {"nil", "none"}:
                        errors.append("Three-line table has grid or missing explicit border suppression.")
                for edge in ("top", "bottom"):
                    item = borders.find(qn("w:" + edge))
                    if item is None or item.get(qn("w:val")) != "single":
                        errors.append("Three-line table is missing a horizontal rule.")
            for cell in table.rows[0].cells:
                if not cell._tc.xpath("./w:tcPr/w:tcBorders/w:bottom[@w:val='single']"):
                    errors.append("Three-line table is missing its header rule.")
    properties = document.core_properties
    if properties.author or properties.last_modified_by:
        errors.append("Personal author metadata is present.")
    text_parts = []
    for paragraph in iter_paragraphs(document):
        text_parts.append(paragraph.text)
        if paragraph.style.style_id == "SFBody":
            ind = paragraph._p.pPr.ind if paragraph._p.pPr is not None else None
            if ind is None or ind.get(qn("w:firstLineChars")) != "200":
                errors.append("Body paragraph must use a two-character first-line indent.")
            for run in paragraph.runs:
                if run.text and (run.font.size is None or abs(run.font.size.pt - template["typography"]["body_pt"]) > 0.01):
                    errors.append("Body font size differs from the selected template.")
        style_xml = paragraph.style.element.xml if paragraph.style is not None else ""
        if "<w:pBdr" in style_xml:
            errors.append("A used paragraph style contains a decorative border.")
        for value in re.findall(r'<w:color[^>]*w:val="([0-9A-Fa-f]{6})"', style_xml):
            if value.upper() != "000000":
                errors.append("A used paragraph style contains non-black color %s." % value.upper())
        for run in paragraph.runs:
            if not run.text:
                continue
            ascii_font = font_value(run, "ascii")
            east_asia_font = font_value(run, "eastAsia")
            if ascii_font != latin_font:
                errors.append("Run does not explicitly use %s: %r" % (latin_font, run.text[:40]))
            if east_asia_font != cjk_font:
                errors.append("Run does not explicitly use %s: %r" % (cjk_font, run.text[:40]))
            color = run.font.color.rgb
            color_value = str(color) if color is not None else None
            if color_value not in (allowed_colors if paragraph.style.style_id == "SFQuestion" else ALLOWED_COLORS):
                errors.append("Non-black run color %s: %r" % (color_value, run.text[:40]))
    content = "\n".join(text_parts)
    blocked, candidates = privacy_errors_and_warnings(content, prohibited_terms, privacy_mode)
    errors.extend("Privacy blocked: " + label for label in blocked)
    warnings.extend("Privacy candidate: " + label for label in candidates)
    with zipfile.ZipFile(path) as archive:
        from lxml import etree
        names = set(archive.namelist())
        if any(name.startswith("word/comments") for name in names):
            errors.append("Review comments are present in the clean output.")
        if "docProps/custom.xml" in names:
            errors.append("Custom document properties are present.")
        for name in names:
            is_story = name == "word/document.xml" or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
            if not is_story:
                continue
            payload = archive.read(name)
            tree = etree.fromstring(payload, etree.XMLParser(resolve_entities=False, no_network=True))
            if next(tree.iter(qn("w:ins")), None) is not None or next(tree.iter(qn("w:del")), None) is not None:
                errors.append("Tracked changes are present in the clean output.")
            math_text = " ".join(node.text or "" for node in tree.iter(qn("m:t")))
            blocked_math, candidates_math = privacy_errors_and_warnings(math_text, prohibited_terms, privacy_mode)
            errors.extend("Privacy blocked in math: " + label for label in blocked_math)
            warnings.extend("Privacy candidate in math: " + label for label in candidates_math)
            for match in re.findall(rb'w:(?:fill|color)="([0-9A-Fa-f]{6})"', payload):
                value = match.decode("ascii").upper()
                if value not in {"000000", "FFFFFF", layout["question_color"]}:
                    errors.append("Non-black/white OOXML color %s in %s." % (value, name))
    errors = sorted(set(errors))
    if not prohibited_terms:
        warnings.append("No task-specific prohibited terms were supplied.")
    return errors, warnings


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path, nargs="+")
    parser.add_argument("--prohibited-term", action="append", default=[])
    parser.add_argument("--template")
    parser.add_argument("--config", type=Path, help="Use the same pack configuration for layout overrides.")
    parser.add_argument("--privacy-mode", choices=("balanced", "strict"), default="balanced")
    args = parser.parse_args()
    config = json.loads(args.config.read_text(encoding="utf-8")) if args.config else {}
    configured = {item["filename"]: item for item in config.get("documents", [])}
    reports = []
    has_errors = False
    for path in args.docx:
        errors, warnings = audit(path, args.prohibited_term, args.template, args.privacy_mode, configured.get(path.name))
        if args.config and path.name not in configured:
            errors.append("Document is not listed in the supplied config.")
        has_errors = has_errors or bool(errors)
        reports.append({"file": path.name, "valid": not errors, "errors": errors, "warnings": warnings})
    print(json.dumps({"valid": not has_errors, "documents": reports}, ensure_ascii=False, indent=2))
    return 1 if has_errors else 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main())
