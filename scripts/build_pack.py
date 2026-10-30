#!/usr/bin/env python3
"""Build a selective, privacy-safe study pack with a deterministic academic DOCX template."""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from pathlib import Path

from common import (
    collect_strings,
    resolve_layout,
    filename_error,
    get_template,
    load_style_catalog,
    privacy_errors_and_warnings,
)

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError as exc:
    raise SystemExit("python-docx is required: python -m pip install -r requirements.txt") from exc


SCENES = {
    "knowledge_point", "exam_prep", "job_interview", "work_onboarding",
    "weekly_progress", "meeting_prep",
}
DOCUMENT_TYPES = {"study_handbook", "teaching_narrative", "practice_guide", "question_bank", "mock_exam"}
REQUIRED_TYPES = {"study_handbook", "teaching_narrative"}
BLOCK_TYPES = {
    "heading", "paragraph", "bullets", "numbered", "table", "code",
    "note", "qa", "checklist", "evidence_table", "page_break", "equation",
}
LATIN_FONT = "Times New Roman"
CJK_FONT = "SimSun"
BLACK = RGBColor(0, 0, 0)
WHITE_HEX = "FFFFFF"
BLACK_HEX = "000000"
CONTENT_WIDTH_DXA = 8300
LABELS = {
    "zh-CN": {
        "scene": "场景", "toc": "目录", "note": "说明",
        "question": "问", "answer": "答", "evidence": "证据账本",
        "evidence_headers": ["编号", "状态", "事实或结论", "来源"],
    },
    "en": {
        "scene": "Scenario", "toc": "Contents", "note": "Note",
        "question": "Q", "answer": "A", "evidence": "Evidence ledger",
        "evidence_headers": ["ID", "Status", "Claim", "Source"],
    },
}
EVIDENCE_STATUSES = {"VERIFIED", "USER_STATED", "INFERRED", "UNKNOWN"}


def validate_block(block, prefix, errors):
    if not isinstance(block, dict):
        errors.append(prefix + " must be an object.")
        return
    block_type = block.get("type")
    if block_type not in BLOCK_TYPES:
        errors.append(prefix + ".type is unsupported.")
        return
    allowed_fields = {
        "heading": {"type", "text", "level"},
        "paragraph": {"type", "text", "bold", "first_indent"},
        "code": {"type", "text"},
        "note": {"type", "text"},
        "bullets": {"type", "items"},
        "numbered": {"type", "items"},
        "checklist": {"type", "items"},
        "table": {"type", "headers", "rows", "widths"},
        "qa": {"type", "question", "answer", "question_blocks", "answer_blocks", "evidence_ids", "origin", "points"},
        "equation": {"type", "omml"},
        "evidence_table": {"type", "ids"},
        "page_break": {"type"},
    }[block_type]
    unsupported = set(block) - allowed_fields
    if unsupported:
        errors.append(prefix + " contains unsupported fields: " + ", ".join(sorted(unsupported)))
    if block_type in {"heading", "paragraph", "code", "note"}:
        if not isinstance(block.get("text"), str) or not block.get("text", "").strip():
            errors.append(prefix + ".text must be a non-empty string.")
    if block_type == "heading":
        if block.get("level", 2) not in {2, 3}:
            errors.append(prefix + ".level must be 2 or 3.")
    if block_type == "paragraph":
        for field in ("bold", "first_indent"):
            if field in block and not isinstance(block[field], bool):
                errors.append(prefix + ".%s must be true or false." % field)
    if block_type in {"bullets", "numbered", "checklist"}:
        items = block.get("items")
        if not isinstance(items, list) or not items or not all(isinstance(item, str) and item.strip() for item in items):
            errors.append(prefix + ".items must be a non-empty list of strings.")
    if block_type == "table":
        headers, rows = block.get("headers"), block.get("rows")
        if not isinstance(headers, list) or not headers or not all(isinstance(item, str) for item in headers):
            errors.append(prefix + ".headers must be a non-empty list of strings.")
        if not isinstance(rows, list) or not all(isinstance(row, list) for row in rows):
            errors.append(prefix + ".rows must be a list of lists.")
        elif isinstance(headers, list) and any(len(row) != len(headers) for row in rows):
            errors.append(prefix + " rows must match the header width.")
        elif any(any(value is None or isinstance(value, (dict, list)) for value in row) for row in rows):
            errors.append(prefix + ".rows cells must be scalar values.")
        widths = block.get("widths")
        if widths is not None:
            if not isinstance(widths, list) or not headers or len(widths) != len(headers):
                errors.append(prefix + ".widths must match the header width.")
            elif not all(isinstance(item, (int, float)) and item > 0 for item in widths):
                errors.append(prefix + ".widths values must be positive numbers.")
    if block_type == "equation":
        try:
            parse_math(block.get("omml"))
        except (ValueError, TypeError) as exc:
            errors.append(prefix + ".omml: " + str(exc))
    if block_type == "qa":
        if block.get("origin", "inferred") not in {"source", "variant", "inferred"}:
            errors.append(prefix + ".origin is unsupported.")
        if "points" in block and (type(block["points"]) not in (int, float) or not 0 < block["points"] < float("inf")):
            errors.append(prefix + ".points must be positive and finite.")
        if "evidence_ids" in block and (not isinstance(block["evidence_ids"], list) or not all(isinstance(x, str) for x in block["evidence_ids"])):
            errors.append(prefix + ".evidence_ids must be a list of strings.")
        for field in ("question_blocks", "answer_blocks"):
            children = block.get(field, [])
            if not isinstance(children, list):
                errors.append(prefix + "." + field + " must be a list.")
                continue
            for i, child in enumerate(children):
                if not isinstance(child, dict) or child.get("type") not in {"paragraph", "equation", "table", "code"}:
                    errors.append(prefix + "." + field + " supports paragraph, equation, table and code only.")
                else:
                    validate_block(child, prefix + "." + field + str(i), errors)
        if not isinstance(block.get("question"), str) or not block.get("question", "").strip():
            errors.append(prefix + ".question must be a non-empty string.")
        if not isinstance(block.get("answer"), str) or not block.get("answer", "").strip():
            errors.append(prefix + ".answer must be a non-empty string.")
    if block_type == "evidence_table":
        ids = block.get("ids")
        if ids is not None and (
            not isinstance(ids, list)
            or not ids
            or not all(isinstance(item, str) and item.strip() for item in ids)
        ):
            errors.append(prefix + ".ids must be a non-empty list of evidence IDs when supplied.")


def validate_config(config):
    errors, warnings = [], []
    if not isinstance(config, dict):
        return ["Configuration root must be an object."], warnings
    if config.get("scene") not in SCENES:
        errors.append("scene must be one of: " + ", ".join(sorted(SCENES)))
    allowed_root = {"scene", "topic", "template", "meta", "selection", "privacy", "evidence", "documents"}
    unsupported_root = set(config) - allowed_root
    if unsupported_root:
        errors.append("configuration contains unsupported fields: " + ", ".join(sorted(unsupported_root)))
    try:
        get_template(config.get("template"))
    except ValueError as exc:
        errors.append(str(exc))
    if not isinstance(config.get("topic"), str) or not config.get("topic", "").strip():
        errors.append("topic must be a non-empty string.")
    meta = config.get("meta", {})
    if not isinstance(meta, dict):
        errors.append("meta must be an object.")
    else:
        for key in ("subtitle", "date", "scene_label", "language"):
            if key in meta and not isinstance(meta[key], str):
                errors.append("meta.%s must be a string." % key)
        if meta.get("language", "zh-CN") not in LABELS:
            errors.append("meta.language must be zh-CN or en.")
        disallowed = set(meta) - {"subtitle", "date", "scene_label", "language"}
        if disallowed:
            errors.append("meta contains unsupported fields: " + ", ".join(sorted(disallowed)))
    selection = config.get("selection")
    if not isinstance(selection, dict):
        errors.append("selection must be an object.")
        selection = {}
    practice_needed = selection.get("practice_needed")
    if not isinstance(practice_needed, bool):
        errors.append("selection.practice_needed must be true or false.")
    if not isinstance(selection.get("practice_reason"), str) or not selection.get("practice_reason", "").strip():
        errors.append("selection.practice_reason must be a non-empty string.")
    output_mode = selection.get("output_mode", "full_pack")
    if output_mode not in {"full_pack", "selected"}:
        errors.append("selection.output_mode must be full_pack or selected.")
    requested_documents = selection.get("requested_documents")
    if output_mode == "selected":
        if (
            not isinstance(requested_documents, list)
            or not requested_documents
            or not all(item in DOCUMENT_TYPES for item in requested_documents)
        ):
            errors.append("selection.requested_documents must list supported document types in selected mode.")
        elif len(set(requested_documents)) != len(requested_documents):
            errors.append("selection.requested_documents must not contain duplicates.")
    elif requested_documents is not None:
        errors.append("selection.requested_documents is only valid in selected mode.")
    unsupported_selection = set(selection) - {
        "output_mode", "requested_documents", "practice_needed", "practice_reason"
    }
    if unsupported_selection:
        errors.append("selection contains unsupported fields: " + ", ".join(sorted(unsupported_selection)))
    privacy = config.get("privacy", {})
    if not isinstance(privacy, dict):
        errors.append("privacy must be an object.")
        privacy = {}
    prohibited = privacy.get("prohibited_terms", [])
    if not isinstance(prohibited, list) or not all(isinstance(term, str) and term.strip() for term in prohibited):
        errors.append("privacy.prohibited_terms must be a list of non-empty strings.")
        prohibited = []
    privacy_mode = privacy.get("mode", "balanced")
    if privacy_mode not in {"balanced", "strict"}:
        errors.append("privacy.mode must be balanced or strict.")
        privacy_mode = "balanced"
    unsupported_privacy = set(privacy) - {"mode", "prohibited_terms"}
    if unsupported_privacy:
        errors.append("privacy contains unsupported fields: " + ", ".join(sorted(unsupported_privacy)))
    evidence = config.get("evidence", [])
    evidence_ids = set()
    if not isinstance(evidence, list):
        errors.append("evidence must be a list.")
        evidence = []
    for index, item in enumerate(evidence):
        prefix = "evidence[%d]" % index
        if not isinstance(item, dict):
            errors.append(prefix + " must be an object.")
            continue
        unsupported = set(item) - {"id", "status", "claim", "source"}
        if unsupported:
            errors.append(prefix + " contains unsupported fields: " + ", ".join(sorted(unsupported)))
        evidence_id = item.get("id")
        if not isinstance(evidence_id, str) or not evidence_id.strip():
            errors.append(prefix + ".id must be a non-empty string.")
        elif evidence_id in evidence_ids:
            errors.append(prefix + ".id is duplicated.")
        else:
            evidence_ids.add(evidence_id)
        if item.get("status") not in EVIDENCE_STATUSES:
            errors.append(prefix + ".status is unsupported.")
        for field in ("claim", "source"):
            if not isinstance(item.get(field), str) or not item.get(field, "").strip():
                errors.append(prefix + ".%s must be a non-empty string." % field)
    documents = config.get("documents")
    if not isinstance(documents, list):
        errors.append("documents must be a list.")
        documents = []
    if output_mode == "full_pack":
        expected_count = 3 if practice_needed is True else 2
        if len(documents) != expected_count:
            errors.append("documents must contain exactly %d document(s) in full_pack mode." % expected_count)
    elif output_mode == "selected" and not documents:
        errors.append("documents must contain at least one document in selected mode.")
    seen = set()
    seen_filenames = set()
    for index, document in enumerate(documents):
        prefix = "documents[%d]" % index
        if not isinstance(document, dict):
            errors.append(prefix + " must be an object.")
            continue
        unsupported_document = set(document) - {"document_type", "filename", "title", "subtitle", "sections", "layout", "duration_minutes", "total_points"}
        if unsupported_document:
            errors.append(prefix + " contains unsupported fields: " + ", ".join(sorted(unsupported_document)))
        doc_type = document.get("document_type")
        if doc_type not in DOCUMENT_TYPES:
            errors.append(prefix + ".document_type is unsupported.")
        elif doc_type in seen:
            errors.append(prefix + ".document_type is duplicated.")
        else:
            seen.add(doc_type)
        filename = document.get("filename")
        filename_problem = filename_error(filename)
        if filename_problem:
            errors.append(prefix + ".filename " + filename_problem)
        elif filename.casefold() in seen_filenames:
            errors.append(prefix + ".filename is duplicated.")
        else:
            seen_filenames.add(filename.casefold())
        for field in ("title", "subtitle"):
            if not isinstance(document.get(field, ""), str):
                errors.append(prefix + ".%s must be a string." % field)
        if not document.get("title", "").strip():
            errors.append(prefix + ".title must be non-empty.")
        try:
            resolve_layout(document)
        except ValueError as exc:
            errors.append(prefix + ": " + str(exc))
        if doc_type == "mock_exam":
            for field in ("duration_minutes", "total_points"):
                if type(document.get(field)) not in (int, float) or not 0 < document[field] < float("inf"):
                    errors.append(prefix + "." + field + " must be positive and finite.")
        document_block_types = []
        question_points = []
        sections = document.get("sections")
        if not isinstance(sections, list) or not sections:
            errors.append(prefix + ".sections must be a non-empty list.")
            sections = []
        for section_index, section in enumerate(sections):
            section_prefix = "%s.sections[%d]" % (prefix, section_index)
            if not isinstance(section, dict):
                errors.append(section_prefix + " must be an object.")
                continue
            unsupported_section = set(section) - {"title", "content"}
            if unsupported_section:
                errors.append(section_prefix + " contains unsupported fields: " + ", ".join(sorted(unsupported_section)))
            if not isinstance(section.get("title"), str) or not section.get("title", "").strip():
                errors.append(section_prefix + ".title must be non-empty.")
            blocks = section.get("content")
            if not isinstance(blocks, list):
                errors.append(section_prefix + ".content must be a list.")
                continue
            for block_index, block in enumerate(blocks):
                if isinstance(block, dict) and block.get("type") == "qa":
                    question_points.append(block.get("points", 0))
                    for evidence_id in block.get("evidence_ids", []) if isinstance(block.get("evidence_ids", []), list) else []:
                        if evidence_id not in evidence_ids:
                            errors.append(section_prefix + " references unknown evidence ID " + str(evidence_id))
                validate_block(block, "%s.content[%d]" % (section_prefix, block_index), errors)
                if isinstance(block, dict) and block.get("type") in BLOCK_TYPES:
                    document_block_types.append(block["type"])
                if isinstance(block, dict) and block.get("type") == "evidence_table":
                    if not evidence:
                        errors.append("%s uses evidence_table but the evidence ledger is empty." % section_prefix)
                    for evidence_id in block.get("ids", []):
                        if evidence_id not in evidence_ids:
                            errors.append("%s references unknown evidence ID %s." % (section_prefix, evidence_id))
        if doc_type == "study_handbook" and "qa" not in document_block_types:
            errors.append(prefix + " study_handbook must include at least one qa block with a question and answer.")
        if doc_type in {"question_bank", "mock_exam"}:
            if not question_points:
                errors.append(prefix + " requires answered qa blocks.")
            if doc_type == "mock_exam" and all(type(x) in (int, float) for x in question_points):
                if any(x <= 0 for x in question_points) or sum(question_points) != document.get("total_points"):
                    errors.append(prefix + " question points must sum to total_points.")
        if doc_type == "teaching_narrative":
            paragraph_count = document_block_types.count("paragraph")
            outline_count = sum(document_block_types.count(item) for item in (
                "bullets", "numbered", "checklist", "table", "evidence_table"
            ))
            if paragraph_count < 2:
                errors.append(prefix + " teaching_narrative must include at least two explanatory paragraphs.")
            if outline_count > max(2, paragraph_count):
                warnings.append(prefix + " teaching_narrative is outline-heavy; prefer connected explanatory prose.")
    if output_mode == "full_pack" and not REQUIRED_TYPES.issubset(seen):
        errors.append("full_pack mode requires study_handbook and teaching_narrative.")
    if output_mode == "full_pack" and practice_needed is True and "practice_guide" not in seen:
        errors.append("practice_needed=true requires practice_guide in full_pack mode.")
    if practice_needed is False and "practice_guide" in seen:
        errors.append("practice_needed=false forbids practice_guide.")
    if output_mode == "selected" and isinstance(requested_documents, list) and set(requested_documents) != seen:
        errors.append("documents must match selection.requested_documents in selected mode.")
    public_content = {key: value for key, value in config.items() if key != "privacy"}
    privacy_errors, privacy_warnings = privacy_errors_and_warnings(
        "\n".join(collect_strings(public_content)), prohibited, privacy_mode
    )
    if privacy_errors:
        errors.append("Privacy scan blocked configuration content: " + ", ".join(privacy_errors))
    if privacy_warnings:
        warnings.append("Privacy candidates require review: " + ", ".join(privacy_warnings))
    if not prohibited:
        warnings.append("privacy.prohibited_terms is empty; add task-specific identifiers when private inputs are used.")
    return errors, warnings


def parse_math(xml):
    from lxml import etree
    if not isinstance(xml, str) or "<!" in xml:
        raise ValueError("Provide OMML without declarations, DTDs or entities.")
    try:
        node = etree.fromstring(xml.encode("utf-8"), etree.XMLParser(resolve_entities=False, no_network=True))
    except etree.XMLSyntaxError as exc:
        raise ValueError("Invalid OMML XML") from exc
    if node.tag != qn("m:oMath"):
        raise ValueError("Root must be m:oMath.")
    for child in node.iter():
        if not isinstance(child.tag, str) or not child.tag.startswith(qn("m:oMath").split("}")[0] + "}"):
            raise ValueError("Only math namespace elements are supported.")
    if not node.findall(".//" + qn("m:t")):
        raise ValueError("Equation must contain mathematical text.")
    required = {"f": ("num", "den"), "sSub": ("e", "sub"), "sSup": ("e", "sup"), "sSubSup": ("e", "sub", "sup")}
    for element in node.iter():
        for tag in required.get(element.tag.split("}")[-1], ()):
            if element.find(qn("m:" + tag)) is None:
                raise ValueError("Incomplete math structure: " + element.tag.split("}")[-1])
        if element.tag == qn("m:m"):
            widths = [len(row.findall(qn("m:e"))) for row in element.findall(qn("m:mr"))]
            if not widths or not widths[0] or len(set(widths)) != 1:
                raise ValueError("Matrix must be non-empty and rectangular.")
    return node


def set_indent(paragraph, enabled, body_pt=12):
    paragraph.paragraph_format.first_line_indent = Pt(2 * body_pt if enabled else 0)
    paragraph._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "200" if enabled else "0")


def color_math_and_text(element, color):
    paragraphs = list(element.iter(qn("w:p")))
    for paragraph in paragraphs:
        props = paragraph.find(qn("w:pPr"))
        if props is None:
            props = OxmlElement("w:pPr")
            paragraph.insert(0, props)
        style = props.find(qn("w:pStyle"))
        if style is None:
            style = OxmlElement("w:pStyle")
            props.insert(0, style)
        style.set(qn("w:val"), "SFQuestion")
    for run in element.iter(qn("w:r")):
        props = run.find(qn("w:rPr"))
        if props is None:
            props = OxmlElement("w:rPr")
            run.insert(0, props)
        node = props.find(qn("w:color"))
        if node is None:
            node = OxmlElement("w:color")
            props.append(node)
        node.set(qn("w:val"), color)
    for run in element.iter(qn("m:r")):
        props = OxmlElement("w:rPr")
        node = OxmlElement("w:color")
        node.set(qn("w:val"), color)
        props.append(node)
        run.insert(0, props)


def apply_table_style(table, mode):
    if mode == "grid":
        return
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    for border in borders:
        edge = border.tag.split("}")[-1]
        border.set(qn("w:val"), "single" if edge in {"top", "bottom"} else "nil")
        border.set(qn("w:sz"), "12")
    for cell in table.rows[0].cells:
        borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        for key, value in (("val", "single"), ("sz", "6"), ("color", "000000")):
            bottom.set(qn("w:" + key), value)
        borders.append(bottom)
        cell._tc.get_or_add_tcPr().append(borders)


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr, value in (("ascii", LATIN_FONT), ("hAnsi", LATIN_FONT), ("cs", LATIN_FONT), ("eastAsia", CJK_FONT)):
        fonts.set(qn("w:" + attr), value)


def set_style_font(style, size, bold=False):
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr, value in (("ascii", LATIN_FONT), ("hAnsi", LATIN_FONT), ("cs", LATIN_FONT), ("eastAsia", CJK_FONT)):
        fonts.set(qn("w:" + attr), value)


def set_cell_margins(cell, value=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        element = margins.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag in ("tblW", "tblInd", "tblLayout"):
        existing = tbl_pr.find(qn("w:" + tag))
        if existing is not None:
            tbl_pr.remove(existing)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BLACK_HEX)
        borders.append(node)
    tbl_pr.append(borders)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        row_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_pr.append(cant_split)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
            if shading is None:
                shading = OxmlElement("w:shd")
                cell._tc.get_or_add_tcPr().append(shading)
            shading.set(qn("w:fill"), WHITE_HEX)
            shading.set(qn("w:val"), "clear")
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_pr.append(repeat_header)


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))


def add_box_border(paragraph):
    p_pr = paragraph._element.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "4")
        node.set(qn("w:color"), BLACK_HEX)
        borders.append(node)
    p_pr.append(borders)


class FormalDocumentBuilder:
    def __init__(self, pack, document_config):
        global LATIN_FONT, CJK_FONT
        self.pack = pack
        self.config = document_config
        self.template_name, self.template = get_template(pack.get("template"))
        LATIN_FONT = self.template["fonts"]["latin"]
        CJK_FONT = self.template["fonts"]["cjk"]
        self.typography = self.template["typography"]
        self.layout = resolve_layout(document_config)
        self.question_number = 0
        language = pack.get("meta", {}).get("language", "zh-CN")
        self.labels = LABELS[language]
        self.doc = Document()
        self._setup_document()

    def _setup_document(self):
        section = self.doc.sections[0]
        margins = self.template["page"]["margins_cm"]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(margins["top"])
        section.bottom_margin = Cm(margins["bottom"])
        section.left_margin = Cm(margins["left"])
        section.right_margin = Cm(margins["right"])
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.5)
        section.different_first_page_header_footer = self.layout["cover"]
        styles = self.doc.styles
        from docx.enum.style import WD_STYLE_TYPE
        question_style = styles.add_style("SFQuestion", WD_STYLE_TYPE.PARAGRAPH)
        question_style.base_style = styles["Normal"]
        body_style = styles.add_style("SFBody", WD_STYLE_TYPE.PARAGRAPH)
        body_style.base_style = styles["Normal"]
        typography = self.typography
        style_specs = {
            "Normal": (typography["body_pt"], False, 0, 6, typography["body_line_spacing"]),
            "Title": (typography["title_pt"], True, 0, 18, 1.0),
            "Subtitle": (typography["subtitle_pt"], False, 0, 12, 1.0),
            "Heading 1": (typography["heading_1_pt"], True, 18, 8, 1.2),
            "Heading 2": (typography["heading_2_pt"], True, 14, 6, 1.2),
            "Heading 3": (typography["heading_3_pt"], True, 10, 4, 1.2),
            "List Bullet": (typography["body_pt"], False, 0, 4, typography["body_line_spacing"]),
            "List Number": (typography["body_pt"], False, 0, 4, typography["body_line_spacing"]),
        }
        style_specs["SFQuestion"] = style_specs["Normal"]
        style_specs["SFBody"] = style_specs["Normal"]
        for name, (size, bold, before, after, spacing) in style_specs.items():
            style = styles[name]
            set_style_font(style, size, bold)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.line_spacing = spacing
            style_p_pr = style.element.get_or_add_pPr()
            for tag in ("pBdr", "shd"):
                inherited_decoration = style_p_pr.find(qn("w:" + tag))
                if inherited_decoration is not None:
                    style_p_pr.remove(inherited_decoration)
        styles["Normal"].paragraph_format.first_line_indent = Pt(0)
        self.doc.core_properties.title = self.config["title"]
        self.doc.core_properties.subject = self.config["document_type"]
        self.doc.core_properties.category = self.template_name
        self.doc.core_properties.author = ""
        self.doc.core_properties.last_modified_by = ""
        self.doc.core_properties.keywords = ""
        self.doc.core_properties.comments = ""
        for tag in ("created", "modified"):
            node = self.doc.core_properties._element.find(qn("dcterms:" + tag))
            if node is not None:
                self.doc.core_properties._element.remove(node)
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header.paragraph_format.space_after = Pt(0)
        if self.layout["header"]:
            run = header.add_run(self.config["title"])
            set_run_font(run, size=9)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(0)
        add_page_field(footer)

    def add_cover(self):
        for _ in range(5):
            self.doc.add_paragraph()
        title = self.doc.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.first_line_indent = None
        title.paragraph_format.space_after = Pt(18)
        run = title.add_run(self.config["title"])
        set_run_font(run, size=self.typography["title_pt"], bold=True)
        subtitle_text = self.config.get("subtitle") or self.pack.get("meta", {}).get("subtitle", "")
        if subtitle_text:
            subtitle = self.doc.add_paragraph(style="Subtitle")
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle.paragraph_format.first_line_indent = None
            subtitle.paragraph_format.space_after = Pt(12)
            run = subtitle.add_run(subtitle_text)
            set_run_font(run, size=self.typography["subtitle_pt"])
        scene_label = self.pack.get("meta", {}).get("scene_label", self.pack["scene"])
        scene = self.doc.add_paragraph()
        scene.alignment = WD_ALIGN_PARAGRAPH.CENTER
        scene.paragraph_format.first_line_indent = None
        scene.paragraph_format.space_before = Pt(24)
        run = scene.add_run(self.labels["scene"] + ": " + scene_label)
        set_run_font(run, size=self.typography["body_pt"])
        date = self.pack.get("meta", {}).get("date", "").strip()
        if date:
            paragraph = self.doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = None
            run = paragraph.add_run(date)
            set_run_font(run, size=self.typography["body_pt"])
        self.doc.add_page_break()

    def add_toc(self):
        heading = self.doc.add_paragraph(style="Heading 1")
        heading.paragraph_format.first_line_indent = None
        run = heading.add_run(self.labels["toc"])
        set_run_font(run, size=self.typography["heading_1_pt"], bold=True)
        for index, section in enumerate(self.config["sections"], 1):
            paragraph = self.doc.add_paragraph()
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.left_indent = Cm(0.5)
            run = paragraph.add_run("%d. %s" % (index, section["title"]))
            set_run_font(run, size=self.typography["body_pt"])
        self.doc.add_page_break()

    def add_heading(self, text, level):
        paragraph = self.doc.add_paragraph(style="Heading %d" % level)
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(text)
        key = {1: "heading_1_pt", 2: "heading_2_pt", 3: "heading_3_pt"}[level]
        set_run_font(run, size=self.typography[key], bold=True)

    def add_paragraph(self, text, bold=False, first_indent=True, size=None):
        paragraph = self.doc.add_paragraph(style="SFBody" if first_indent else "Normal")
        if first_indent:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_indent(paragraph, first_indent, self.typography["body_pt"])
        run = paragraph.add_run(text)
        set_run_font(run, size=size or self.typography["body_pt"], bold=bold)
        return paragraph

    def add_list(self, items, numbered=False):
        style = "List Number" if numbered else "List Bullet"
        for item in items:
            paragraph = self.doc.add_paragraph(style=style)
            paragraph.paragraph_format.first_line_indent = None
            run = paragraph.add_run(item)
            set_run_font(run, size=self.typography["body_pt"])

    def add_table(self, block):
        headers, rows = block["headers"], block["rows"]
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        weights = block.get("widths") or [1] * len(headers)
        total = float(sum(weights))
        widths = [round(CONTENT_WIDTH_DXA * float(weight) / total) for weight in weights]
        widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
        set_table_geometry(table, widths)
        apply_table_style(table, self.layout["table_style"])
        for row_index, values in enumerate([headers] + rows):
            for column_index, value in enumerate(values):
                cell = table.rows[row_index].cells[column_index]
                cell.text = ""
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.first_line_indent = None
                paragraph.paragraph_format.space_after = Pt(0)
                if column_index == 0 or len(str(value)) <= 12:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(str(value))
                set_run_font(run, size=10.5, bold=row_index == 0)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)

    def add_code(self, text):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.left_indent = Cm(0.4)
        paragraph.paragraph_format.right_indent = Cm(0.4)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15
        add_box_border(paragraph)
        run = paragraph.add_run(text)
        set_run_font(run, size=self.typography["code_pt"])

    def add_note(self, text):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.left_indent = Cm(0.4)
        paragraph.paragraph_format.right_indent = Cm(0.4)
        label = paragraph.add_run(self.labels["note"] + ": ")
        set_run_font(label, size=self.typography["note_pt"], bold=True)
        run = paragraph.add_run(text)
        set_run_font(run, size=self.typography["note_pt"])

    def add_equation(self, omml):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_indent(paragraph, False)
        paragraph._p.append(parse_math(omml))

    def add_qa(self, block, answer_only=False):
        self.question_number += 1
        if not answer_only:
            start = len(self.doc._element.body)
            paragraph = self.add_paragraph(str(self.question_number) + ". " + block["question"], first_indent=False)
            paragraph.paragraph_format.keep_with_next = True
            if "points" in block:
                set_run_font(paragraph.add_run(" (%s)" % block["points"]))
            for child in block.get("question_blocks", []):
                self.process_block(child)
            for element in list(self.doc._element.body)[start-1:-1]:
                color_math_and_text(element, self.layout["question_color"])
        if self.layout["answer_mode"] == "adjacent" or answer_only:
            prefix = (str(self.question_number) + ". " if answer_only else "") + self.labels["answer"] + ": "
            self.add_paragraph(prefix + block["answer"])
            for child in block.get("answer_blocks", []):
                self.process_block(child)

    def add_evidence_table(self, ids=None):
        evidence = self.pack.get("evidence", [])
        if ids:
            selected = set(ids)
            evidence = [item for item in evidence if item["id"] in selected]
        block = {
            "headers": self.labels["evidence_headers"],
            "rows": [[item["id"], item["status"], item["claim"], item["source"]] for item in evidence],
            "widths": [0.8, 1.2, 3.2, 2.0],
        }
        self.add_table(block)

    def process_block(self, block):
        block_type = block["type"]
        if block_type == "heading":
            self.add_heading(block["text"], block.get("level", 2))
        elif block_type == "paragraph":
            self.add_paragraph(block["text"], block.get("bold", False), block.get("first_indent", True))
        elif block_type == "bullets":
            self.add_list(block["items"])
        elif block_type == "numbered":
            self.add_list(block["items"], numbered=True)
        elif block_type == "checklist":
            self.add_list(["[ ] " + item for item in block["items"]])
        elif block_type == "evidence_table":
            self.add_evidence_table(block.get("ids"))
        elif block_type == "table":
            self.add_table(block)
        elif block_type == "code":
            self.add_code(block["text"])
        elif block_type == "note":
            self.add_note(block["text"])
        elif block_type == "qa":
            self.add_qa(block)
        elif block_type == "equation":
            self.add_equation(block["omml"])
        elif block_type == "page_break":
            self.doc.add_page_break()

    def build(self):
        if self.layout["cover"]:
            self.add_cover()
        else:
            title = self.doc.add_paragraph(style="Title")
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(title.add_run(self.config["title"]), size=self.typography["title_pt"], bold=True)
        if self.config["document_type"] == "mock_exam":
            self.add_paragraph(("时长：%s 分钟　满分：%s 分" if self.pack.get("meta", {}).get("language", "zh-CN") == "zh-CN" else "%s min | %s points") % (self.config["duration_minutes"], self.config["total_points"]), first_indent=False)
        if self.layout["toc"]:
            self.add_toc()
        for section in self.config["sections"]:
            self.add_heading(section["title"], 1)
            for block in section["content"]:
                self.process_block(block)
        if self.layout["answer_mode"] == "separate":
            self.doc.add_page_break()
            self.add_heading(self.labels["answer"], 1)
            self.question_number = 0
            for section in self.config["sections"]:
                for block in section["content"]:
                    if block["type"] == "qa":
                        self.add_qa(block, answer_only=True)
        return self.doc


def load_json(path):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise SystemExit("Configuration not found: %s" % path) from exc
    except json.JSONDecodeError as exc:
        raise SystemExit("Invalid JSON at line %d, column %d: %s" % (exc.lineno, exc.colno, exc.msg)) from exc


def extract_document_text(document):
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(parts)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--check-only", action="store_true")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--list-templates", action="store_true")
    args = parser.parse_args()
    if args.list_templates:
        catalog = load_style_catalog()
        print(json.dumps({
            "default": catalog["default"],
            "templates": [
                {"name": name, "label": item.get("label", name), "fonts": item["fonts"]}
                for name, item in catalog["templates"].items()
            ],
        }, ensure_ascii=False, indent=2))
        return 0
    if args.config is None:
        parser.error("--config is required unless --list-templates is used")
    config = load_json(args.config)
    errors, warnings = validate_config(config)
    result = {"valid": not errors, "errors": errors, "warnings": warnings,
              "scene": config.get("scene") if isinstance(config, dict) else None}
    if errors:
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 2
    if args.check_only:
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    if args.output_dir is None:
        parser.error("--output-dir is required unless --check-only is used")
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    targets = [output_dir / item["filename"] for item in config["documents"]]
    existing = [path.name for path in targets if path.exists()]
    if existing and not args.overwrite:
        raise SystemExit("Output exists; choose a new directory or add --overwrite: " + ", ".join(existing))
    prohibited = config.get("privacy", {}).get("prohibited_terms", [])
    privacy_mode = config.get("privacy", {}).get("mode", "balanced")
    built = []
    with tempfile.TemporaryDirectory(prefix="study-fast-", dir=str(output_dir)) as temporary_dir:
        temp_root = Path(temporary_dir)
        for item, target in zip(config["documents"], targets):
            temporary = temp_root / target.name
            document = FormalDocumentBuilder(config, item).build()
            document.save(temporary)
            reopened = Document(temporary)
            blocked, candidates = privacy_errors_and_warnings(
                extract_document_text(reopened), prohibited, privacy_mode
            )
            if blocked:
                raise SystemExit("Privacy scan blocked generated document %s: %s" % (target.name, ", ".join(blocked)))
            if candidates:
                result.setdefault("warnings", []).append(
                    "Privacy candidates in %s require review: %s" % (target.name, ", ".join(candidates))
                )
            built.append((temporary, target))
        for temporary, target in built:
            os.replace(temporary, target)
    result["outputs"] = [{"filename": path.name, "bytes": path.stat().st_size} for path in targets]
    result["practice_needed"] = config["selection"]["practice_needed"]
    result["practice_reason"] = config["selection"]["practice_reason"]
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main())
