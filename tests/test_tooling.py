from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
EXAMPLE = ROOT / "assets" / "pack-config.example.json"


def run_script(name, *arguments):
    return subprocess.run(
        [sys.executable, str(SCRIPTS / name), *map(str, arguments)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )


class ToolingTests(unittest.TestCase):
    def test_two_document_pack_validates_builds_and_audits(self):
        checked = run_script("build_pack.py", "--config", EXAMPLE, "--check-only")
        self.assertEqual(checked.returncode, 0, checked.stdout)
        self.assertTrue(json.loads(checked.stdout)["valid"])
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            built = run_script("build_pack.py", "--config", EXAMPLE, "--output-dir", output)
            self.assertEqual(built.returncode, 0, built.stdout)
            files = sorted(output.glob("*.docx"))
            self.assertEqual(len(files), 2)
            audited = run_script("audit_docx.py", *files)
            self.assertEqual(audited.returncode, 0, audited.stdout)

    def test_practice_policy_requires_and_generates_third_document(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = json.loads(EXAMPLE.read_text(encoding="utf-8"))
            config["selection"] = {
                "practice_needed": True,
                "practice_reason": "现场编程要求提供可运行代码和失败测试。",
            }
            invalid = root / "invalid.json"
            invalid.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            result = run_script("build_pack.py", "--config", invalid, "--check-only")
            self.assertEqual(result.returncode, 2, result.stdout)
            practice = {
                "document_type": "practice_guide",
                "filename": "03-practice-guide.docx",
                "title": "技术实操指南",
                "subtitle": "代码、测试与追问",
                "sections": [{
                    "title": "最小练习",
                    "content": [
                        {"type": "paragraph", "text": "本练习用于验证基本推理，不声称来自任何私人项目。"},
                        {"type": "code", "text": "def identity(value):\n    return value"},
                        {"type": "numbered", "items": ["正常输入", "边界输入", "失败输入"]},
                    ],
                }],
            }
            config["documents"].append(practice)
            valid = root / "valid.json"
            valid.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            output = root / "output"
            result = run_script("build_pack.py", "--config", valid, "--output-dir", output)
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertEqual(len(list(output.glob("*.docx"))), 3)

    def test_privacy_scan_blocks_identifiers(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = json.loads(EXAMPLE.read_text(encoding="utf-8"))
            config["documents"][0]["sections"][0]["content"][1]["text"] += " 联系方式 13812345678"
            path = root / "private.json"
            path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            result = run_script("build_pack.py", "--config", path, "--check-only")
            self.assertEqual(result.returncode, 2, result.stdout)
            self.assertIn("Privacy scan", result.stdout)

    def test_balanced_privacy_warns_and_placeholder_is_allowed(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = json.loads(EXAMPLE.read_text(encoding="utf-8"))
            text = " 示例邮箱 learner@example.com，配置 token=YOUR_TOKEN"
            config["documents"][0]["sections"][0]["content"][1]["text"] += text
            path = root / "balanced.json"
            path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            result = run_script("build_pack.py", "--config", path, "--check-only")
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("Privacy candidates require review", result.stdout)

    def test_strict_privacy_blocks_warning_candidates(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = json.loads(EXAMPLE.read_text(encoding="utf-8"))
            config["privacy"]["mode"] = "strict"
            config["documents"][0]["sections"][0]["content"][1]["text"] += " 示例邮箱 learner@example.com"
            path = root / "strict.json"
            path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            result = run_script("build_pack.py", "--config", path, "--check-only")
            self.assertEqual(result.returncode, 2, result.stdout)

    def test_selected_mode_builds_one_document(self):
        example = ROOT / "examples" / "selected-handbook.json"
        checked = run_script("build_pack.py", "--config", example, "--check-only")
        self.assertEqual(checked.returncode, 0, checked.stdout)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            built = run_script("build_pack.py", "--config", example, "--output-dir", output)
            self.assertEqual(built.returncode, 0, built.stdout)
            files = list(output.glob("*.docx"))
            self.assertEqual(len(files), 1)
            audited = run_script("audit_docx.py", *files)
            self.assertEqual(audited.returncode, 0, audited.stdout)

    def test_duplicate_and_reserved_filenames_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = json.loads(EXAMPLE.read_text(encoding="utf-8"))
            config["documents"][1]["filename"] = config["documents"][0]["filename"].upper()
            duplicate = root / "duplicate.json"
            duplicate.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            result = run_script("build_pack.py", "--config", duplicate, "--check-only")
            self.assertEqual(result.returncode, 2, result.stdout)
            self.assertIn("filename is duplicated", result.stdout)
            config["documents"][1]["filename"] = "CON.docx"
            reserved = root / "reserved.json"
            reserved.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            result = run_script("build_pack.py", "--config", reserved, "--check-only")
            self.assertEqual(result.returncode, 2, result.stdout)
            self.assertIn("reserved Windows device name", result.stdout)

    def test_unknown_evidence_reference_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = json.loads(EXAMPLE.read_text(encoding="utf-8"))
            config["documents"][0]["sections"][0]["content"][0]["ids"] = ["MISSING"]
            path = root / "evidence.json"
            path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            result = run_script("build_pack.py", "--config", path, "--check-only")
            self.assertEqual(result.returncode, 2, result.stdout)
            self.assertIn("unknown evidence ID", result.stdout)

    def test_unknown_configuration_field_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = json.loads(EXAMPLE.read_text(encoding="utf-8"))
            config["documents"][0]["sections"][0]["content"][1]["frist_indent"] = False
            path = root / "typo.json"
            path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            result = run_script("build_pack.py", "--config", path, "--check-only")
            self.assertEqual(result.returncode, 2, result.stdout)
            self.assertIn("unsupported fields", result.stdout)

    def test_core_document_learning_contract_is_enforced(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = json.loads(EXAMPLE.read_text(encoding="utf-8"))
            handbook = config["documents"][0]
            for section in handbook["sections"]:
                section["content"] = [block for block in section["content"] if block["type"] != "qa"]
            path = root / "missing-answer.json"
            path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            result = run_script("build_pack.py", "--config", path, "--check-only")
            self.assertEqual(result.returncode, 2, result.stdout)
            self.assertIn("study_handbook must include at least one qa block", result.stdout)

            config = json.loads(EXAMPLE.read_text(encoding="utf-8"))
            narrative = config["documents"][1]
            for section in narrative["sections"]:
                section["content"] = [block for block in section["content"] if block["type"] != "paragraph"]
            path = root / "missing-prose.json"
            path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            result = run_script("build_pack.py", "--config", path, "--check-only")
            self.assertEqual(result.returncode, 2, result.stdout)
            self.assertIn("teaching_narrative must include at least two explanatory paragraphs", result.stdout)

    def test_template_catalog_is_available(self):
        result = run_script("build_pack.py", "--list-templates")
        self.assertEqual(result.returncode, 0, result.stdout)
        payload = json.loads(result.stdout)
        names = {item["name"] for item in payload["templates"]}
        self.assertEqual(names, {"formal_zh_cn", "formal_cross_platform"})

    def test_evidence_checklist_and_table_pagination_markup(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            result = run_script("build_pack.py", "--config", EXAMPLE, "--output-dir", output)
            self.assertEqual(result.returncode, 0, result.stdout)
            path = output / "01-study-handbook.docx"
            document = Document(path)
            self.assertTrue(document.sections[0].different_first_page_header_footer)
            self.assertTrue(any("[ ] " in paragraph.text for paragraph in document.paragraphs))
            self.assertTrue(any("E01" in cell.text for table in document.tables for row in table.rows for cell in row.cells))
            with zipfile.ZipFile(path) as archive:
                document_xml = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("w:tblHeader", document_xml)
            self.assertIn("w:cantSplit", document_xml)

    def test_real_secret_is_blocked_but_placeholder_is_not(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = json.loads(EXAMPLE.read_text(encoding="utf-8"))
            target = config["documents"][0]["sections"][0]["content"][1]
            target["text"] += " api_key=super-secret-value"
            path = root / "secret.json"
            path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            result = run_script("build_pack.py", "--config", path, "--check-only")
            self.assertEqual(result.returncode, 2, result.stdout)
            self.assertIn("secret_assignment", result.stdout)

    def test_fixed_template_fonts_page_and_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            result = run_script("build_pack.py", "--config", EXAMPLE, "--output-dir", output)
            self.assertEqual(result.returncode, 0, result.stdout)
            document = Document(output / "01-study-handbook.docx")
            section = document.sections[0]
            self.assertAlmostEqual(section.page_width.cm, 21.0, places=1)
            self.assertAlmostEqual(section.page_height.cm, 29.7, places=1)
            self.assertEqual(document.core_properties.author, "")
            self.assertEqual(document.core_properties.category, "formal_zh_cn")
            title_paragraph = next(paragraph for paragraph in document.paragraphs if paragraph.text == "缓存一致性快速学习手册")
            self.assertNotIn("<w:pBdr", title_paragraph._element.xml)
            runs = [run for paragraph in document.paragraphs for run in paragraph.runs if run.text]
            self.assertTrue(runs)
            for run in runs:
                fonts = run._element.rPr.rFonts
                self.assertEqual(fonts.get(qn("w:ascii")), "Times New Roman")
                self.assertEqual(fonts.get(qn("w:eastAsia")), "SimSun")
                self.assertEqual(str(run.font.color.rgb), "000000")

    def test_overwrite_guard(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            first = run_script("build_pack.py", "--config", EXAMPLE, "--output-dir", output)
            self.assertEqual(first.returncode, 0, first.stdout)
            second = run_script("build_pack.py", "--config", EXAMPLE, "--output-dir", output)
            self.assertNotEqual(second.returncode, 0)
            self.assertIn("Output exists", second.stdout)

    def test_repository_self_check(self):
        result = run_script("self_check.py")
        self.assertEqual(result.returncode, 0, result.stdout)
        self.assertEqual(json.loads(result.stdout)["errors"], [])


if __name__ == "__main__":
    unittest.main()
