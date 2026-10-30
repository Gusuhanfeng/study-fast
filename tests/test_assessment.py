import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from build_pack import FormalDocumentBuilder, validate_config, parse_math
from audit_docx import audit

class AssessmentTests(unittest.TestCase):
    def config(self, name='question-bank'):
        return json.loads((ROOT/'examples'/f'{name}.json').read_text(encoding='utf-8'))

    def test_bank_colors_indent_table_and_native_math(self):
        c=self.config(); self.assertFalse(validate_config(c)[0])
        d=FormalDocumentBuilder(c,c['documents'][0]).build()
        question=next(p for p in d.paragraphs if p.text.startswith('1.'))
        self.assertEqual(str(question.runs[0].font.color.rgb),'1F4E79')
        answer=next(p for p in d.paragraphs if p.text.startswith('答:'))
        self.assertEqual(answer._p.pPr.ind.get(qn('w:firstLineChars')),'200')
        self.assertEqual(str(answer.runs[0].font.color.rgb),'000000')
        self.assertTrue(d._element.xpath('.//m:f'))
        self.assertFalse(any(p.text for p in d.sections[0].header.paragraphs))
        self.assertEqual(str(d.tables[0].cell(0,0).paragraphs[0].runs[0].font.color.rgb),'1F4E79')
        with tempfile.TemporaryDirectory() as t:
            path=Path(t)/'bank.docx';d.save(path)
            self.assertFalse(audit(path,[])[0])
            answer.runs[0].font.color.rgb=RGBColor.from_string('1F4E79');d.save(path)
            self.assertTrue(audit(path,[])[0], 'Blue answers must fail role-aware audit')

    def test_exam_answer_section_and_score_failure(self):
        c=self.config('mock-exam');self.assertFalse(validate_config(c)[0])
        d=FormalDocumentBuilder(c,c['documents'][0]).build()
        text=[p.text for p in d.paragraphs]
        self.assertLess(text.index('答'),next(i for i,p in enumerate(text) if p.startswith('1. 答:')))
        c['documents'][0]['total_points']=11
        self.assertTrue(validate_config(c)[0])

    def test_layout_override_and_header_audit(self):
        c=self.config();cfg=c['documents'][0]
        cfg['layout']={'header':True,'table_style':'grid','question_color':'000000'}
        d=FormalDocumentBuilder(c,cfg).build()
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'bank.docx';d.save(p)
            self.assertFalse(audit(p,[],document_config=cfg)[0])
            self.assertTrue(audit(p,[])[0])

    def test_missing_answer_bad_evidence_and_bad_layout(self):
        for mutate in [lambda c:c['documents'][0]['sections'][0]['content'][0].update(answer=''),
                       lambda c:c['documents'][0]['sections'][0]['content'][0].update(evidence_ids=['missing']),
                       lambda c:c['documents'][0].update(layout={'header':'no'}),
                       lambda c:c['documents'][0]['sections'][0]['content'][0].update(question_blocks=[{'type':'qa'}])]:
            c=self.config();mutate(c);self.assertTrue(validate_config(c)[0])

    def test_math_rejects_entities_and_non_math_elements(self):
        for xml in ['<!DOCTYPE x><x/>','<x/>','<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><x/></m:oMath>']:
            with self.assertRaises(ValueError):parse_math(xml)

    def test_matrix_and_subscript_remain_structured(self):
        xml='<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:m><m:mr><m:e><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>1</m:t></m:r></m:sub></m:sSub></m:e><m:e><m:r><m:t>2</m:t></m:r></m:e></m:mr></m:m></m:oMath>'
        c=self.config();c['documents'][0]['sections'][0]['content'][0]['answer_blocks']=[{'type':'equation','omml':xml}]
        self.assertFalse(validate_config(c)[0]);d=FormalDocumentBuilder(c,c['documents'][0]).build()
        self.assertEqual(len(d._element.xpath('.//m:m/m:mr/m:e')),2)
        self.assertEqual(len(d._element.xpath('.//m:sSub')),1)

    def test_broken_body_indent_is_detected(self):
        c=self.config();d=FormalDocumentBuilder(c,c['documents'][0]).build()
        answer=next(p for p in d.paragraphs if p.text.startswith('答:'))
        answer._p.pPr.ind.set(qn('w:firstLineChars'),'0')
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'bank.docx';d.save(p)
            self.assertTrue(any('two-character' in e for e in audit(p,[])[0]))

    def test_incomplete_fraction_is_rejected(self):
        with self.assertRaises(ValueError):
            parse_math('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num></m:f></m:oMath>')

    def test_internal_provenance_is_not_printed(self):
        c=self.config();d=FormalDocumentBuilder(c,c['documents'][0]).build()
        text=' '.join(p.text for p in d.paragraphs)
        self.assertNotIn('INFERRED',text);self.assertNotIn('E1',text)

if __name__=='__main__':unittest.main()
